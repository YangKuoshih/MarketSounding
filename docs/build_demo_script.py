"""Generate demo-script.docx -- MarketBuzz hackathon presentation script + reference glossary."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MID_BLUE  = RGBColor(0x2E, 0x5F, 0x8C)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAY      = RGBColor(0x55, 0x55, 0x55)
GREEN     = RGBColor(0x1A, 0x6B, 0x3A)
AMBER     = RGBColor(0x7A, 0x4F, 0x00)

doc = Document()

section = doc.sections[0]
section.top_margin    = Inches(0.9)
section.bottom_margin = Inches(0.9)
section.left_margin   = Inches(1.1)
section.right_margin  = Inches(1.1)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    r = p.add_run(text)
    r.font.name  = 'Calibri'
    r.font.size  = Pt(15)
    r.font.bold  = True
    r.font.color.rgb = DARK_BLUE

def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), '1B3A5C')
    pBdr.append(left)
    pPr.append(pBdr)
    r = p.add_run(text)
    r.font.name  = 'Calibri'
    r.font.size  = Pt(12)
    r.font.bold  = True
    r.font.color.rgb = DARK_BLUE

def body(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    _inline(p, text)

def action(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.left_indent  = Inches(0.2)
    r = p.add_run(text)
    r.font.name    = 'Calibri'
    r.font.size    = Pt(10)
    r.font.italic  = True
    r.font.color.rgb = GRAY

def say(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Inches(0.35)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F5FB')
    pPr.append(shd)
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '10')
    left.set(qn('w:color'), '2E5F8C')
    pBdr.append(left)
    pPr.append(pBdr)
    r = p.add_run('"' + text + '"')
    r.font.name  = 'Calibri'
    r.font.size  = Pt(11)
    r.font.color.rgb = DARK_BLUE

def type_box(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.left_indent  = Inches(0.35)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E8F5E9')
    pPr.append(shd)
    rb = p.add_run('TYPE:  ')
    rb.font.name  = 'Calibri'
    rb.font.size  = Pt(10)
    rb.font.bold  = True
    rb.font.color.rgb = GREEN
    rc = p.add_run(text)
    rc.font.name  = 'Consolas'
    rc.font.size  = Pt(10)
    rc.font.color.rgb = GREEN

def tip(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    p.paragraph_format.left_indent  = Inches(0.2)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'FFF8E1')
    pPr.append(shd)
    rb = p.add_run('TIP  ')
    rb.font.name = 'Calibri'
    rb.font.size = Pt(10)
    rb.font.bold = True
    rb.font.color.rgb = AMBER
    rc = p.add_run(text)
    rc.font.name = 'Calibri'
    rc.font.size = Pt(10)
    rc.font.color.rgb = AMBER

def divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'BDD0E0')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for run in hdr_cells[i].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = WHITE
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1B3A5C')
        tcPr.append(shd)
    for ri, row_data in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, cell_text in enumerate(row_data):
            cells[ci].text = cell_text
            for run in cells[ci].paragraphs[0].runs:
                run.font.name = 'Calibri'
                run.font.size = Pt(10)
            if ri % 2 == 1:
                tc = cells[ci]._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'EEF4FB')
                tcPr.append(shd)
    if col_widths:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Inches(col_widths[i])
    doc.add_paragraph()

def _inline(para, text):
    tokens = re.split(r'(\*\*[^*]+\*\*)', text)
    for tok in tokens:
        if tok.startswith('**') and tok.endswith('**'):
            r = para.add_run(tok[2:-2])
            r.font.bold = True
            r.font.color.rgb = DARK_BLUE
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
        else:
            r = para.add_run(tok)
            r.font.name = 'Calibri'
            r.font.size = Pt(11)


# ===========================================================================
# COVER
# ===========================================================================

cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before = Pt(72)
cover.paragraph_format.space_after  = Pt(6)
r = cover.add_run('MarketBuzz')
r.font.name  = 'Calibri'
r.font.size  = Pt(36)
r.font.bold  = True
r.font.color.rgb = DARK_BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(6)
r = sub.add_run('Hackathon Presentation Script')
r.font.name   = 'Calibri'
r.font.size   = Pt(15)
r.font.italic = True
r.font.color.rgb = MID_BLUE

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub2.paragraph_format.space_after = Pt(4)
r = sub2.add_run('6 - 7 minutes  |  Reference copy for presenter')
r.font.name  = 'Calibri'
r.font.size  = Pt(11)
r.font.color.rgb = GRAY

doc.add_page_break()


# ===========================================================================
# SETUP CHECKLIST
# ===========================================================================

h1('Before You Present -- Setup Checklist')
body('Complete these steps before walking to the front of the room:')
for item in [
    "Log in as  tony / password123  at localhost:3000",
    "Console dashboard open at localhost:3000/console",
    "Jarrett widget visible in the bottom-right corner",
    "Pre-run at least one simulation so the knowledge graph has data",
    "Confirm the backend is running on port 3001 (dev.bat)",
]:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(item)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)

divider()


# ===========================================================================
# SEGMENT 1
# ===========================================================================

h2('Segment 1 -- Opening Hook  [0:00 - 0:40]')
action('[On screen: Landing page  localhost:3000 -- scroll slowly as you speak]')
say("In macro markets, understanding how primary dealers are positioned matters. "
    "Analysts track dealer sentiment through published commentary, research notes, "
    "and periodic surveys -- but those are snapshots. By the time they are published, "
    "the market has often already moved.")
say("MarketBuzz gives analysts an on-demand simulation of dealer sentiment -- how Goldman, "
    "JPMorgan, Morgan Stanley, Citi, and BofA would likely react to any market event, "
    "right now, before the next note hits your inbox.")
say("It is not a chatbot. It is a structured AI debate between five distinct dealer personas "
    "-- and it completes in under two minutes.")

divider()


# ===========================================================================
# SEGMENT 2
# ===========================================================================

h2('Segment 2 -- Dashboard Overview  [0:40 - 1:10]')
action('[Navigate to localhost:3000/console]')
say("Here is the console. You can launch a new simulation by typing any market topic "
    "directly into the search bar, or browse prior runs to track how simulated dealer "
    "sentiment has evolved over time.")
action('[Point to the recent simulation cards]')
say("Each card shows a consensus score and trajectory at a glance. "
    "We will explain exactly what those mean.")

divider()


# ===========================================================================
# SEGMENT 3
# ===========================================================================

h2('Segment 3 -- The H/D Score: What the Simulation Shows  [1:10 - 2:00]')
action('[Click "New Simulation" in the nav -- localhost:3000/console/sounding/new]')
say("The core output is a Hawkish/Dovish score -- a number from negative one to positive "
    "one for each dealer.")
say("Negative means dovish: the dealer favors rate cuts, looser monetary policy, more "
    "accommodation for growth. Positive means hawkish: they favor rate hikes, tighter "
    "conditions, inflation control. Zero is neutral.")
say("Each dealer has a distinct institutional lean built into their persona. Goldman and "
    "Citi lean slightly dovish -- they weight growth risk and the lagged effects of policy. "
    "Morgan Stanley and BofA lean more hawkish -- they weight inflation persistence and "
    "balance sheet risk. JPMorgan sits near neutral, anchored to labor market data.")
say("When the simulation runs you see all five scores side by side. Tightly clustered means "
    "high consensus -- the Street agrees. Spread out means genuine disagreement -- which is "
    "strategically the more interesting signal, because it tells you positioning can diverge.")
action('[Type into the topic field:]')
type_box('Bank of Japan surprise rate hike -- 25bps above consensus')
action('[Set rounds to 3, leave crisis off, click Launch]')

divider()


# ===========================================================================
# SEGMENT 4
# ===========================================================================

h2('Segment 4 -- Multi-Agent Rounds: How the Debate Works  [2:00 - 3:15]')
action('[Simulation is running -- point to the progress indicators while speaking]')
say("Now let me explain what multi-agent actually means here -- because this is not five "
    "chatbots answering the same question.")
say("First, a research agent runs parallel live web searches and synthesizes a current event "
    "brief. That becomes the shared context.")
say("Round one: all five dealer agents run simultaneously -- each is a separate Claude Sonnet "
    "model call on AWS Bedrock, loaded with a deep persona. Each agent has a documented "
    "analytical voice, known blind spots, institutional bias, and signature phrases that make "
    "their output sound unmistakably like their desk. They produce an initial Hawkish/Dovish "
    "score, a rate path view, a balance sheet view, and a key quote.")
say("Round two is where it gets interesting. Each dealer reads the other four positions and "
    "responds -- just like an actual roundtable. They can shift their score, but only if a "
    "peer raises a point that is compelling through their own framework. Goldman is not going "
    "to suddenly sound like Morgan Stanley. Each desk holds its character.")
say("That delta -- who moved, who held firm, who influenced whom -- is the signal.")
say("Optionally, you can inject a crisis event mid-simulation: a sudden credit shock, a "
    "surprise policy reversal. All five dealers re-evaluate under pressure. You see who "
    "pivots and who stays anchored.")
tip("The ~90-second wait is your pipeline explainer window. Use every second of it.")

divider()


# ===========================================================================
# SEGMENT 5
# ===========================================================================

h2('Segment 5 -- Reading the Simulation Report  [3:15 - 4:30]')
action('[Simulation completes -- results view loads automatically]')
say("Let me walk through what you are looking at here, section by section.")

action('[Point to the two summary cards at the very top -- Consensus Score and Overall Trajectory]')
say("At the top: two headline numbers. The consensus score tells you how aligned the Street "
    "is -- a high percentage means dealers converged to a similar view. The trajectory tells "
    "you which direction they landed and whether they moved there over rounds or arrived "
    "immediately. This is your one-line read before you go deeper.")

action('[Point to the H/D spectrum bar showing all 5 dealers, directly below the summary cards]')
say("Below that is the H/D spectrum -- all five dealers plotted on the same axis at their "
    "final round scores. This is the most important visual. You want to see the spread. "
    "If Goldman is at negative 0.5 and Morgan Stanley is at positive 0.6, that is a "
    "1.1-point spread -- the Street is genuinely divided. For a BoJ hike you might expect "
    "exactly that: Goldman hedging on global growth spillovers while Morgan Stanley argues "
    "the inflation transmission into US rates is real.")

action('[Point to the individual dealer cards]')
say("Each dealer card gives you the full breakdown. Rate path view -- what they think the "
    "Fed does next. Balance sheet view -- their position on quantitative tightening or "
    "easing. Risk asset view -- how they see equities and credit reacting. Key concerns -- "
    "the specific risks their desk is watching. And a key quote -- one sentence in their "
    "documented institutional voice. You should be able to read that quote and immediately "
    "know which desk wrote it, without seeing the name.")

action('[Point to the position evolution chart]')
say("The position evolution chart is where you see the debate play out over time. Each line "
    "is one dealer across rounds. A flat line means they were anchored -- high conviction, "
    "not moved by peers. A line that shifts toward another dealer tells you who drove "
    "consensus. If Goldman moved toward JPMorgan in round two and Citi followed in round "
    "three, JPMorgan was the anchor of that debate. That influence dynamic is what no "
    "traditional survey captures.")

action('[Point to the reasoning section under each dealer card]')
say("Finally, the reasoning section gives you the full analytical narrative -- two to three "
    "paragraphs in each dealer's documented voice, referencing their specific models and "
    "frameworks. This is the most useful section for analysts: it explains not just the "
    "position, but the logic behind it.")

divider()


# ===========================================================================
# SEGMENT 6
# ===========================================================================

h2('Segment 6 -- Jarrett: Your AI Research Guide  [4:30 - 5:30]')
action('[Click the Jarrett widget in the bottom-right corner]')
say("This is Jarrett -- our AI market intelligence assistant. He is not a generic chatbot. "
    "He is purpose-built for this platform and serves three distinct functions.")

say("First: he explains simulation results in plain English. After a sim completes you can "
    "ask Jarrett what the output means -- he knows the scoring system, the dealer personas, "
    "and the context of the event. He will tell you which dealer drove consensus, what a "
    "specific H/D score implies for positioning, and whether a low consensus score is a "
    "warning signal or just noise.")
action('[Type into Jarrett:]')
type_box("What does a hawkish lean from Morgan Stanley on a BoJ hike mean for US rate positioning?")
action('[Wait briefly for reply, then continue]')

say("Second: he is an agentic orchestrator. You do not need to fill out a form to run a "
    "simulation. You just describe what you want.")
action('[Type into Jarrett:]')
type_box("Run a simulation on the Fed holding rates at the June FOMC")
say("Jarrett interprets the request, configures the parameters -- topic, number of rounds, "
    "crisis toggle -- outputs a structured simulation spec, fills the form automatically, "
    "and launches the pipeline. No dropdowns. No menus. You just talk to him.")

say("Third: on the full chat page Jarrett can route your question directly to any individual "
    "dealer. Ask what Goldman thinks about duration risk -- he answers in Goldman's documented "
    "analytical voice, referencing their models and frameworks. Ask what Citi thinks -- "
    "Citi's voice. One interface to all five desks, available any time.")

divider()


# ===========================================================================
# SEGMENT 7
# ===========================================================================

h2('Segment 7 -- Knowledge Graph  [5:30 - 6:20]')
action('[Click "Graph" in the nav -- localhost:3000/console/graph]')
say("Every simulation we run contributes to a persistent knowledge graph. This is not a "
    "one-time visualization -- it accumulates automatically across every run.")

say("The nodes represent three types of things: dealers, market events or topics, and concern "
    "categories -- things like inflation persistence, financial conditions tightening, or EM "
    "contagion risk. The edges represent relationships: which dealer raised which concern, "
    "which topic connects to which risk category, and which dealers positions converged.")

action('[Point to the graph -- highlight different node types if visible]')
say("What you can read from this graph: which concerns appear most frequently -- those are "
    "the persistent risks the Street keeps coming back to. Which dealers cluster together -- "
    "those are the desks that tend to agree. And which topics are most connected -- those "
    "are the events that activate the broadest set of market concerns.")

action('[Open Jarrett widget, type:]')
type_box("Show me the shortest path between Goldman Sachs and inflation risk")
action('[Point to the highlighted path in the D3 graph]')
say("Jarrett queries the graph directly. Here he is running a shortest-path analysis -- "
    "showing which nodes link Goldman to inflation risk, and through what intermediaries. "
    "You can also ask for centrality analysis to find the most influential nodes, filter by "
    "concern type, or pull a subgraph of just the dealers to see influence networks.")

say("Every query is saved to history and replayable. For an analyst running weekly "
    "simulations, this graph becomes a living institutional memory -- showing how the "
    "Street's concerns and alignments have shifted across events over the entire research cycle.")

divider()


# ===========================================================================
# SEGMENT 8
# ===========================================================================

h2('Segment 8 -- Close  [6:20 - 7:00]')
action('[Navigate back to localhost:3000/console]')
say("On the technical side: Next.js, Node/Express, five parallel Claude Sonnet 4.6 dealer "
    "agents on AWS Bedrock, Tavily for live research, D3 for the knowledge graph, Jarrett "
    "as the agentic interface layer. All custom-built. No third-party orchestration framework.")
say("But what you have actually seen today is a compression of a workflow that historically "
    "took days. An analyst would wait for the event note, source commentary from each desk, "
    "cross-reference the survey cycle, and piece together where the Street stood. "
    "By the time that synthesis was ready, the window had often already closed.")
say("MarketBuzz compresses that to under two minutes -- not by summarising existing "
    "commentary, but by simulating the debate itself. Five distinct analytical voices, "
    "reacting to the same event, responding to each other, arriving at positions that "
    "reflect each desk's documented framework under peer pressure. The dynamic, not "
    "just the snapshot.")
say("The most valuable moment in market intelligence is the one before the Street "
    "publishes. MarketBuzz is the tool that lives in that window -- and it gets sharper "
    "with every simulation that runs.")
say("Thank you.")

divider()


# ===========================================================================
# TIMING TABLE
# ===========================================================================

h1('Timing Guide')
add_table(
    ['Segment', 'Duration', 'Key Action'],
    [
        ['1 -- Opening Hook',          '0:40', 'Landing page scroll'],
        ['2 -- Dashboard',             '0:30', 'Console overview'],
        ['3 -- H/D Score + Launch',    '0:50', 'New Sim page, launch'],
        ['4 -- Multi-agent Rounds',    '1:15', 'Pipeline explainer while sim runs'],
        ['5 -- Reading the Report',    '1:15', 'Walk each section of the results view'],
        ['6 -- Jarrett',               '1:00', 'Widget -- 3 functions demo'],
        ['7 -- Knowledge Graph',       '0:50', 'Graph + Jarrett query'],
        ['8 -- Close',                 '0:40', 'Stack + value prop + close'],
        ['TOTAL',                      '~7:00', ''],
    ],
    col_widths=[2.3, 1.0, 3.0]
)

divider()


# ===========================================================================
# PRESENTER TIPS
# ===========================================================================

h1('Presenter Tips')
body("**Non-technical judges:** Lead every segment with the analyst use case. They need to "
     "understand it replaces waiting for the next survey cycle -- not that it runs on Bedrock.")
body("**AI-technical judges:** Mention the persona calibration (defaultBias, enforced H/D "
     "spread of 0.6 across all five agents), the parallel Promise.all per round, Jarrett's "
     "agentic tool-block pattern, and the BFS graph traversal in the D3 layer.")
body("**If the sim runs slow:** Use the wait to explain the pipeline in Segment 4. "
     "The 90-second run time is a feature -- it shows real AI inference happening live.")
body("**Pre-warm:** Run a sim right before presenting. It primes the Bedrock endpoint "
     "and ensures the graph has visible data for Segment 7.")

divider()
doc.add_page_break()


# ===========================================================================
# GLOSSARY
# ===========================================================================

h1('Glossary & Reference')
body('Use this section during judge Q&A.')


# --- H/D Score ---

h2('Hawkish / Dovish Score (H/D Score)')
body('The core metric in every simulation. Each dealer receives a score from -1.0 to +1.0 after each round.')
add_table(
    ['Score Range', 'Label', 'What It Means'],
    [
        ['-1.0 to -0.5', 'Strongly Dovish',
         'Favors significant rate cuts; prioritizes growth and employment over inflation control'],
        ['-0.5 to -0.1', 'Mildly Dovish',
         'Leans toward accommodation; cautious on overtightening'],
        ['-0.1 to +0.1', 'Neutral',
         'Balanced view; data-dependent, no strong policy lean'],
        ['+0.1 to +0.5', 'Mildly Hawkish',
         'Leans toward tightening; watchful on inflation persistence'],
        ['+0.5 to +1.0', 'Strongly Hawkish',
         'Favors rate hikes or sustained tight policy; inflation control priority'],
    ],
    col_widths=[1.3, 1.4, 3.6]
)


# --- Dealer Leans ---

h2('Dealer Institutional Leans')
add_table(
    ['Dealer', 'Default Bias', 'Analytical Anchor'],
    [
        ['Goldman Sachs (GS)',     '-0.2  Slightly Dovish',
         'Quantitative models; weights lagged policy effects and growth risk'],
        ['JP Morgan (JPM)',         '0.0   Neutral',
         'Labor market focus; data-dependent, pragmatic'],
        ['Morgan Stanley (MS)',     '+0.3  Mildly Hawkish',
         'Tail-risk framing; inflation persistence; market underpricing risk'],
        ['Citigroup (Citi)',        '-0.1  Slightly Dovish',
         'Global macro; EM contagion; financial conditions sensitivity'],
        ['Bank of America (BofA)', '+0.2  Mildly Hawkish',
         'Consumer credit health; balance sheet policy; card data signals'],
    ],
    col_widths=[1.8, 1.7, 2.8]
)


# --- Simulation Metrics ---

h2('Simulation Metrics')
add_table(
    ['Term', 'Definition'],
    [
        ['H/D Score',
         'Hawkish/Dovish score per dealer per round (-1 to +1). Primary output of each agent invocation.'],
        ['Consensus Score',
         'How tightly the five dealers final H/D scores cluster. 100% = full agreement. ~40% = Street is genuinely split.'],
        ['Position Shift',
         'How much a dealer score changed from the prior round after reading peers. Max +/-0.3 per round unless warranted.'],
        ['Trajectory',
         'Average H/D score across all five dealers plotted per round. Shows whether group sentiment converged, diverged, or reversed.'],
        ['Influence Edge',
         'Graph edge from Dealer A to B means B score shifted toward A view in a round. Thicker = larger shift.'],
        ['Crisis Injection',
         'Optional mid-simulation shock event that forces all five dealers to re-evaluate from their institutional framework.'],
        ['Initial Round',
         'Round 1: each dealer produces their first independent reaction. No peer context yet.'],
        ['Peer Response Round',
         'Round 2+: each dealer reads all peers prior positions and updates their own. Core of the debate dynamics.'],
        ['Crisis Re-evaluation',
         'Final round when crisis is enabled: dealers re-evaluate after the injected shock. Largest position shifts typically here.'],
    ],
    col_widths=[1.8, 4.5]
)


# --- How to Read the Report ---

h2('How to Read a Simulation Report -- Section by Section')
add_table(
    ['Section', 'What It Shows', 'How to Interpret It'],
    [
        ['Consensus Score',
         'A 0-100% measure of how tightly the five dealers final H/D scores cluster.',
         'High (>70%): the Street agrees -- strong directional signal. Low (<40%): genuine divergence -- market may be mis-pricing.'],
        ['Overall Trajectory',
         'Average H/D score across all five dealers per round, plotted as a line.',
         'Moving dovish: group shifted toward easier policy. Flat: consensus formed early. Volatile: genuine disagreement resolved over rounds.'],
        ['H/D Spectrum Bar',
         'All five dealers plotted on the -1 to +1 axis at their final round scores.',
         'Wide spread (>0.8) = Street is divided. Narrow = consensus. Lone outliers carry the most strategic signal.'],
        ['Dealer Cards',
         'Per-dealer: H/D score, rate path view, balance sheet view, risk asset view, key concerns, key quote.',
         'Rate path = what they think the Fed does next. Key quote should be instantly recognizable as that desk.'],
        ['Position Evolution Chart',
         'Each dealer H/D score plotted across all rounds as individual lines.',
         'Flat line = anchored, high conviction. Converging lines = consensus forming. Dealer everyone moves toward is the anchor.'],
        ['Reasoning / Narrative',
         'Each dealer full analytical explanation in their documented institutional voice.',
         'Explains WHY they landed at their score. Should reference specific models -- GS Financial Conditions Index, JPM labor indicators, etc.'],
        ['Influenced By',
         'Which dealers caused a given dealer position to shift.',
         'Only in round 2+. Shows who moved whom. Dealer with most outbound influence edges is the consensus driver for that event.'],
        ['Final Positions',
         'Each dealer settled view after the full debate.',
         'The actionable output. Simulated Street positioning after hearing each other -- closer to how real dealer consensus forms.'],
    ],
    col_widths=[1.5, 2.1, 2.7]
)


# --- Jarrett Capabilities ---

h2('Jarrett -- Capabilities Reference')
add_table(
    ['Function', 'How to Use It', 'Example'],
    [
        ['Explain results',
         'After a simulation, ask Jarrett what the output means in plain English.',
         '"What does a 40% consensus score mean for how I should read the Street?"'],
        ['Interpret a dealer position',
         'Ask Jarrett what a specific dealer score or reasoning implies.',
         '"What does a hawkish lean from Morgan Stanley on a BoJ hike mean for US rates?"'],
        ['Launch a simulation',
         'Describe what you want to simulate. Jarrett fills the form and launches automatically.',
         '"Run a 3-round simulation on the Fed holding rates at the June FOMC"'],
        ['Talk to a specific dealer',
         'On the /chat page, ask about a named dealer. Jarrett routes to that dealer persona.',
         '"What does Goldman Sachs think about duration risk given today macro data?"'],
        ['Query the knowledge graph',
         'Ask Jarrett to find connections, paths, or clusters in the graph.',
         '"Show me the shortest path between Goldman Sachs and inflation risk"'],
        ['Compare dealers',
         'Ask for a cross-desk comparison on any topic.',
         '"How do Goldman and Morgan Stanley differ on rate cut timing?"'],
    ],
    col_widths=[1.5, 2.2, 2.6]
)


# --- Knowledge Graph ---

h2('Knowledge Graph -- How to Use It')
add_table(
    ['Node Type', 'What It Represents', 'How to Use It'],
    [
        ['Dealer node',
         'One of the 5 primary dealers. Size reflects simulation count and avg H/D score.',
         'Find dealers that cluster together -- those desks tend to agree across events.'],
        ['Topic node',
         'A simulated market event (e.g., BoJ surprise rate hike).',
         'Highly connected topic nodes activate many dealer concerns simultaneously.'],
        ['Concern node',
         'A macro concern raised by a dealer (e.g., inflation persistence).',
         'High-frequency concern nodes appear across many simulations -- persistent Street risks.'],
        ['Influence edge',
         'Arrow from Dealer A to B when B score shifted toward A in a round.',
         'Thick edges = strong influence. Most outbound edges = consensus driver.'],
        ['Topic edge',
         'Connects a topic node to a dealer node.',
         'Shows which dealers have the most exposure to a given market event.'],
        ['Concern edge',
         'Connects a dealer to a concern category they raised.',
         'Reveals each desk persistent analytical focus across all simulations.'],
    ],
    col_widths=[1.3, 2.5, 2.5]
)
add_table(
    ['Query Type', 'What It Does', 'When to Use It'],
    [
        ['Shortest Path',
         'Finds the minimum-hop connection between two nodes.',
         'Understand how two seemingly unrelated concepts link through dealer networks.'],
        ['Centrality',
         'Ranks nodes by how many connections they have.',
         'Find the most influential topics or concerns across all simulations.'],
        ['Filter by Type',
         'Shows only dealer, topic, or concern nodes.',
         'Declutter the graph and focus on one layer at a time.'],
        ['Filter by Concern',
         'Highlights all dealers and topics connected to a concern.',
         'Find every dealer and event associated with a specific risk theme.'],
        ['Highlight Node',
         'Dims everything except a node and its direct connections.',
         'Focus on one dealer or topic without losing graph context.'],
        ['Subgraph',
         'Extracts a subset of nodes and edges into a focused view.',
         'Compare two dealers concern profiles side by side.'],
    ],
    col_widths=[1.3, 2.5, 2.5]
)


# --- Tech Stack ---

h2('Tech Stack (Quick Reference)')
add_table(
    ['Layer', 'Technology', 'Purpose'],
    [
        ['Frontend',       'Next.js, TypeScript, Tailwind CSS, Framer Motion',
         'React app, App Router, animations'],
        ['Charts / Graph', 'Highcharts, D3.js',
         'Position charts, force-directed knowledge graph'],
        ['Backend',        'Node.js, Express, TypeScript',
         'API orchestration, auth, sim management'],
        ['AI Models',      'Claude Sonnet 4.6 -- dealer agents + synthesis',
         'AWS Bedrock, parallel dealer invocations'],
        ['AI Models',      'Claude Haiku 4.5 -- Jarrett chat + query gen',
         'AWS Bedrock, low-latency interactive use'],
        ['Search',         'Tavily Search API',
         'Parallel live web research for event briefs'],
        ['Auth',           'JWT + SHA-256',
         'Token-based auth, password hashing'],
        ['Persistence',    'JSON file (local) / DynamoDB (prod)',
         'Users, simulations, graph nodes/edges'],
        ['Agentic Layer',  'Custom TypeScript orchestrator',
         'No third-party framework -- built from scratch'],
    ],
    col_widths=[1.3, 2.5, 2.5]
)


# --- Judge Questions ---

h2('Common Judge Questions')
add_table(
    ['Question', 'Short Answer'],
    [
        ['What agentic framework?',
         "Custom TypeScript orchestration over direct Bedrock calls -- no LangChain or AutoGen. "
         "We own the loop: parallel fan-out per round, sequential rounds, tool-block pattern for Jarrett."],
        ['What is your backend?',
         "Node.js Express API in TypeScript. Structured as Lambda handlers that run locally as a "
         "monolith and deploy to AWS serverless unchanged. Bedrock for AI, Tavily for search, JWT auth."],
        ['Why Claude over GPT-4?',
         "Claude's instruction-following fidelity for structured JSON output is superior. "
         "Our 11-field dealer schema requires exact JSON every invocation -- Claude's lower "
         "hallucination rate on structured outputs reduced retry rates significantly."],
        ['Why AWS Bedrock?',
         "Compliance posture: Bedrock keeps data within AWS infrastructure -- critical for a "
         "Federal Reserve context. Also enables cross-region inference for burst load when "
         "five agents run in parallel."],
        ['How accurate are the personas?',
         "Personas are built from publicly documented dealer house styles, key voices, and known "
         "analytical frameworks. They are a simulation, not official commentary -- but calibrated "
         "to reproduce each desk known biases and blind spots."],
        ['What would production look like?',
         "Same Lambda handlers deploy to AWS directly. DynamoDB replaces the JSON file. API "
         "Gateway fronts the Express routes. The local server is a development convenience, not a rewrite."],
    ],
    col_widths=[1.8, 4.5]
)


# ===========================================================================
# SAVE
# ===========================================================================

output_path = 'd:/Users/tony.yang/Documents/MarketSounding/docs/demo-script.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
